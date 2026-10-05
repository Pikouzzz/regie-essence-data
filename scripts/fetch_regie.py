"""
Récupère le dernier export Excel de la Régie de l'énergie (regieessencequebec.ca)
et le convertit en stations.json pour l'application.

Utilisation :
  python scripts/fetch_regie.py                 # mode normal (GitHub Actions)
  python scripts/fetch_regie.py --file x.xlsx   # test avec un fichier local
"""
import io, json, re, sys, time
from urllib.parse import urlparse
from datetime import datetime, timedelta, timezone
from pathlib import Path

import requests
from openpyxl import load_workbook
from shapely.geometry import Point, shape

BASE      = "https://regieessencequebec.ca"
XLSX_URL  = BASE + "/data/stations-{ts}.xlsx"
OUT       = Path(__file__).resolve().parent.parent / "stations.json"
RIVES     = Path(__file__).resolve().parent.parent / "data" / "rives.geojson"
CORRECTIONS = Path(__file__).resolve().parent.parent / "data" / "rives_corrections.json"
QC_BOITE  = (44.9, 63.5, -80.0, -56.0)   # lat min, lat max, lng min, lng max : hors de ça = coordonnées erronées
MIN_STATIONS = 1000            # en dessous, on considère l'export invalide
PRIX_MIN, PRIX_MAX = 50, 400   # ¢/L — hors de cette plage = donnée aberrante
TAILLE_MAX = 20 * 1024 * 1024   # 20 Mo : un export normal fait ~300 Ko
HEADERS   = {"User-Agent": "regie-essence-data (github.com/Pikouzzz/regie-essence-data)"}

# Colonnes attendues dans l'Excel (repérées par leur nom, pas leur position)
COLS = {
    "nom": "Nom", "banniere": "Bannière", "adresse": "Adresse", "region": "Région",
    "cp": "Code Postal", "lat": "Latitude", "lng": "Longitude",
    "regulier": "Prix Régulier", "super": "Prix Super", "diesel": "Prix Diesel",
}

session = requests.Session()
session.headers.update(HEADERS)


def log(*a):
    print(*a, flush=True)


DELAI = (5, 10)            # secondes : connexion, lecture
MAX_AGE_TOLERE_H = 3       # si la Régie est injoignable, on tolère des données de moins de 3 h sans alerter
DUREE_MAX_S = 240          # le robot abandonne après 4 min au lieu de bloquer 15 min


class SiteInjoignable(Exception):
    pass


def existe(url):
    """'ok' si l'URL répond 200, 'absent' si 404/403…, 'erreur' si le site ne répond pas."""
    try:
        with session.get(url, stream=True, timeout=DELAI) as r:
            return "ok" if r.status_code == 200 else "absent"
    except requests.RequestException:
        return "erreur"


def trouver_dernier_export():
    """Retourne (url, horodatage 'AAAAMMJJHHMMSS') du plus récent export disponible."""
    debut = time.monotonic()
    # 1) Chercher le lien directement dans la page d'accueil
    try:
        html = session.get(BASE + "/", timeout=DELAI).text
        liens = sorted(set(re.findall(r"stations-(\d{14})\.xlsx", html)))
        if liens:
            ts = liens[-1]
            log(f"Lien trouvé dans la page : {ts}")
            return XLSX_URL.format(ts=ts), ts
    except requests.RequestException as e:
        log(f"Page d'accueil inaccessible ({e}), on sonde les URL.")

    # 2) Sonder les tranches de 5 min récentes (UTC), secondes 00 à 09
    now = datetime.now(timezone.utc)
    slot = now.replace(second=0, microsecond=0, minute=now.minute - now.minute % 5)
    erreurs_de_suite = 0
    for k in range(9):                         # jusqu'à 40 min en arrière
        base = slot - timedelta(minutes=5 * k)
        for sec in range(10):
            if time.monotonic() - debut > DUREE_MAX_S:
                raise SiteInjoignable("délai maximal dépassé pendant le sondage")
            ts = (base + timedelta(seconds=sec)).strftime("%Y%m%d%H%M%S")
            etat = existe(XLSX_URL.format(ts=ts))
            if etat == "ok":
                log(f"Export trouvé par sondage : {ts}")
                return XLSX_URL.format(ts=ts), ts
            erreurs_de_suite = erreurs_de_suite + 1 if etat == "erreur" else 0
            if erreurs_de_suite >= 3:
                raise SiteInjoignable("le site ne répond pas (3 délais dépassés de suite)")
    return None, None


def age_donnees_h():
    """Âge en heures des données actuellement publiées (None si inconnu)."""
    try:
        ts = json.loads(OUT.read_text(encoding="utf-8"))["ts"]
        dt = datetime.strptime(ts, "%Y-%m-%dT%H:%M:%SZ").replace(tzinfo=timezone.utc)
        return (datetime.now(timezone.utc) - dt).total_seconds() / 3600
    except Exception:
        return None


def abandon_tolere(raison):
    """La Régie est injoignable : pas d'alerte si nos données sont encore récentes."""
    age = age_donnees_h()
    if age is not None and age < MAX_AGE_TOLERE_H:
        log(f"::warning::Régie injoignable ({raison}). Données actuelles conservées (âge : {age:.1f} h).")
        sys.exit(0)
    raise SystemExit(f"Régie injoignable ({raison}) et données trop vieilles ({'inconnu' if age is None else f'{age:.1f} h'}).")


def prix(v):
    """'179.9¢' -> 179.9 ; 'N/D' / vide / aberrant -> None"""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        p = float(v)
    else:
        m = re.search(r"\d+(?:[.,]\d+)?", str(v))
        if not m:
            return None
        p = float(m.group().replace(",", "."))
    return round(p, 1) if PRIX_MIN <= p <= PRIX_MAX else None


def num(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def stable_id(banniere, adresse, lat, lng, seen):
    """Même hash FNV-1a que dans l'app (sur les unités UTF-16, comme charCodeAt)."""
    s = f"{banniere}|{adresse}|{'' if lat is None else f'{lat:.4f}'}|{'' if lng is None else f'{lng:.4f}'}"
    h = 2166136261
    b = s.encode("utf-16-le")
    for i in range(0, len(b), 2):
        h ^= b[i] | (b[i + 1] << 8)
        h = (h * 16777619) & 0xFFFFFFFF
    base = ""
    n = h
    chars = "0123456789abcdefghijklmnopqrstuvwxyz"
    while True:
        n, r = divmod(n, 36)
        base = chars[r] + base
        if n == 0:
            break
    seen[base] = seen.get(base, 0) + 1
    return base if seen[base] == 1 else f"{base}-{seen[base]}"


# Rive de chaque station (N = nord du fleuve, S = sud, M = Îles-de-la-Madeleine)
# Polygones simplifiés tirés de Natural Earth (domaine public), précision ≈ 1-2 km :
# une station en bord de fleuve est rattachée à la rive la plus proche.
_rives = None
def rive(lat, lng):
    global _rives
    if lat is None or lng is None:
        return None
    if _rives is None:
        gj = json.loads(RIVES.read_text(encoding="utf-8"))
        _rives = [(f["properties"]["rive"], shape(f["geometry"])) for f in gj["features"]]
    pt = Point(lng, lat)
    for code, geom in _rives:
        if geom.contains(pt):
            return code
    return min(_rives, key=lambda r: r[1].distance(pt))[0]


def coord_valide(lat, lng):
    return (lat is not None and lng is not None
            and QC_BOITE[0] <= lat <= QC_BOITE[1] and QC_BOITE[2] <= lng <= QC_BOITE[3])


def lire_excel(contenu):
    wb = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = ws.iter_rows(values_only=True)
    entete = [str(c).strip() if c is not None else "" for c in next(rows)]
    manquantes = [v for v in COLS.values() if v not in entete]
    if manquantes:
        raise SystemExit(f"Colonnes introuvables dans l'Excel : {manquantes} (en-tête lu : {entete})")
    idx = {k: entete.index(v) for k, v in COLS.items()}

    stations, seen = [], {}
    for r in rows:
        if not r or all(c is None for c in r):
            continue
        g = lambda k: r[idx[k]]
        # Texte nettoyé : caractères de contrôle retirés, longueur limitée
        txt = lambda k, n=150: re.sub(r"[\x00-\x1f\x7f]", " ", str(g(k))).strip()[:n] if g(k) is not None else ""
        lat, lng = num(g("lat")), num(g("lng"))
        if not coord_valide(lat, lng):
            lat = lng = None
        banniere = txt("banniere")
        stations.append([
            stable_id(banniere, txt("adresse"), lat, lng, seen),
            txt("nom"), banniere, txt("adresse"), txt("region"), txt("cp"),
            lat, lng, prix(g("regulier")), prix(g("super")), prix(g("diesel")),
            rive(lat, lng),
        ])
    return stations


def main():
    args = sys.argv[1:]
    if args[:1] == ["--file"]:
        chemin = Path(args[1])
        contenu = chemin.read_bytes()
        m = re.search(r"(\d{14})", chemin.name)
        ts, url = (m.group(1) if m else datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")), str(chemin)
    else:
        try:
            url, ts = trouver_dernier_export()
        except SiteInjoignable as e:
            abandon_tolere(str(e))
        if not url:
            abandon_tolere("aucun export trouvé ; le format d'URL a peut-être changé")
        # Rien à faire si on a déjà cet export
        if OUT.exists():
            try:
                if json.loads(OUT.read_text(encoding="utf-8")).get("source", "").endswith(ts + ".xlsx"):
                    log("Déjà à jour.")
                    return
            except json.JSONDecodeError:
                pass
        for essai in range(3):
            try:
                r = session.get(url, timeout=(5, 60))
                r.raise_for_status()
                # Sécurité : on n'accepte que le domaine de la Régie (même après redirection) et une taille raisonnable
                if not (urlparse(r.url).hostname or "").endswith("regieessencequebec.ca"):
                    raise SystemExit(f"Redirection vers un domaine inattendu : {r.url}")
                if len(r.content) > TAILLE_MAX:
                    raise SystemExit("Fichier anormalement gros : abandon.")
                contenu = r.content
                break
            except requests.RequestException as e:
                log(f"Téléchargement échoué ({e}), nouvel essai…")
                time.sleep(5)
        else:
            abandon_tolere("téléchargement impossible")

    stations = lire_excel(contenu)
    # Corrections manuelles de rive (stations en bord de fleuve mal classées)
    if CORRECTIONS.exists():
        corr = json.loads(CORRECTIONS.read_text(encoding="utf-8"))
        for st in stations:
            if st[0] in corr:
                st[-1] = corr[st[0]]
    if len(stations) < MIN_STATIONS:
        raise SystemExit(f"Seulement {len(stations)} stations : export jugé invalide, on garde l'ancien fichier.")

    dt = datetime.strptime(ts, "%Y%m%d%H%M%S").replace(tzinfo=timezone.utc)
    data = {
        "version": 1,
        "ts": dt.strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": url,
        "count": len(stations),
        "fields": ["id", "nom", "banniere", "adresse", "region", "cp", "lat", "lng", "regulier", "super", "diesel", "rive"],
        "stations": stations,
    }
    OUT.write_text(json.dumps(data, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    log(f"OK : {len(stations)} stations, export du {data['ts']} → {OUT.name}")


if __name__ == "__main__":
    main()
